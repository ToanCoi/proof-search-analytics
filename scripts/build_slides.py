"""Sinh slides/slides.pptx từ results/ (số liệu đọc từ file, dùng chung hàm với build_report.py).

    python scripts/build_slides.py
"""
from __future__ import annotations

import json
import os
import sys

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

sys.path.insert(0, os.path.dirname(__file__))
from build_report import FIG, FINAL_TAU, RAW, live_rescue, runs, share, sweep_of  # noqa: E402

NAVY, ACCENT, GREY, LIGHT = (RGBColor(0x1F, 0x3A, 0x5F), RGBColor(0xC0, 0x39, 0x2B),
                             RGBColor(0x55, 0x5B, 0x66), RGBColor(0xEE, 0xF2, 0xF7))
W, H = Inches(13.333), Inches(7.5)


def _text(slide, x, y, w, h, lines, size=18, color=GREY, bold_first=False, align=None):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        bold = False
        if isinstance(line, tuple):
            line, bold = line
        p.text = line
        p.space_after = Pt(8)
        if align:
            p.alignment = align
        for r in p.runs:
            r.font.size, r.font.color.rgb, r.font.name = Pt(size), color, "Calibri"
            r.font.bold = bold or (bold_first and i == 0)
    return tb


def slide(prs, title, kicker=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.18), H)
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    if kicker:
        _text(s, Inches(0.6), Inches(0.3), Inches(12), Inches(0.4), [kicker], 13, ACCENT)
    _text(s, Inches(0.6), Inches(0.6), Inches(12), Inches(0.9), [title], 30, NAVY, bold_first=True)
    n = len(prs.slides)
    _text(s, Inches(12.3), Inches(7.0), Inches(0.8), Inches(0.4), [str(n)], 11, GREY, align=PP_ALIGN.RIGHT)
    return s


def bullets(s, items, x=Inches(0.6), y=Inches(1.7), w=Inches(12), h=Inches(5), size=20):
    return _text(s, x, y, w, h, [("• " + i if not isinstance(i, tuple) else ("• " + i[0], i[1])) for i in items], size)


def picture(s, path, x, y, w=None, h=None):
    s.shapes.add_picture(path, x, y, width=w, height=h)


def table(s, header, rows, x, y, w, row_h=Inches(0.42), size=14, highlight=()):
    t = s.shapes.add_table(len(rows) + 1, len(header), x, y, w, row_h * (len(rows) + 1)).table
    for i, hd in enumerate(header):
        c = t.cell(0, i)
        c.text = hd
        c.fill.solid()
        c.fill.fore_color.rgb = NAVY
        for p in c.text_frame.paragraphs:
            for r in p.runs:
                r.font.size, r.font.bold, r.font.color.rgb = Pt(size), True, RGBColor(255, 255, 255)
    for k, row in enumerate(rows, 1):
        for i, v in enumerate(row):
            c = t.cell(k, i)
            c.text = str(v)
            c.fill.solid()
            c.fill.fore_color.rgb = RGBColor(0xFD, 0xED, 0xEC) if k - 1 in highlight else \
                (LIGHT if k % 2 else RGBColor(255, 255, 255))
            for p in c.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size, r.font.color.rgb = Pt(size), GREY
                    r.font.bold = k - 1 in highlight
    return t


def build(out="slides/slides.pptx"):
    rs = runs()
    collect = [r for r in rs if r["kind"] == "collect"]
    summ = pd.read_csv(f"{FIG}/search_summary.csv").set_index("run_name")
    scale = pd.read_csv(f"{FIG}/scale_aggregate.csv")
    pm = json.load(open(f"{RAW}/predictor_metrics.json", encoding="utf-8"))
    pm_uni = json.load(open(f"{RAW}/predictor_metrics_unigram.json", encoding="utf-8"))
    sw_path, sw = sweep_of("M")
    m_run = os.path.basename(sw_path)[len("prune_sweep_"):-4]
    cause = pd.read_csv(f"{FIG}/pruning_losses_M_by_cause.csv").set_index("tau")
    bM, fin = summ.loc["repl-baseline-M"], summ.loc["repl-prune-M-tau001"]
    save = 1 - fin.avg_states_expanded_median / bM.avg_states_expanded_median

    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H

    # 1 tiêu đề
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    bg.line.fill.background()
    _text(s, Inches(0.9), Inches(1.0), Inches(11), Inches(0.5), ["INT4418 Big Data · Nhóm 7 · AI4Math"], 18,
          RGBColor(0xC9, 0xD6, 0xE8))
    _text(s, Inches(0.9), Inches(2.0), Inches(11.5), Inches(2.2),
          ["Proof-Search Analytics", "Trace lake cho tìm kiếm chứng minh Lean 4: phân tích, dự đoán tactic, cắt nhánh"],
          40, RGBColor(255, 255, 255), bold_first=True)
    _text(s, Inches(0.9), Inches(5.6), Inches(11), Inches(0.8),
          ["[Họ tên 1] · [Họ tên 2] · [Họ tên 3] · [Họ tên 4]"], 18, RGBColor(0xC9, 0xD6, 0xE8))

    # 2 kết quả chính trước
    s = slide(prs, "Kết quả chính", "Tóm tắt")
    cards = [(f"−{save * 100:.0f}%", f"states expanded ở τ = {FINAL_TAU:g}\n(live Lean, budget M)"),
             (f"{fin.proof_success_rate_median:.2f}", f"proof rate giữ nguyên\n(baseline {bM.proof_success_rate_median:.2f})"),
             (f"{pm['tfidf_sgd']['top1']:.2f}", f"top-1 tactic TF-IDF\n(most-frequent {pm['most_frequent']['top1']:.2f})"),
             (f"{scale.events.max() / 1e6:.1f}M", "event tổng hợp bằng Spark\n(lake nhân bản lớn nhất)")]
    for i, (big, small) in enumerate(cards):
        x = Inches(0.6 + i * 3.1)
        box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.0), Inches(2.8), Inches(2.6))
        box.fill.solid()
        box.fill.fore_color.rgb = LIGHT
        box.line.fill.background()
        _text(s, x, Inches(2.3), Inches(2.8), Inches(1), [big], 44, ACCENT if i == 0 else NAVY, True, PP_ALIGN.CENTER)
        _text(s, x + Inches(0.15), Inches(3.5), Inches(2.5), Inches(1), small.split("\n"), 15, GREY, align=PP_ALIGN.CENTER)
    _text(s, Inches(0.6), Inches(5.2), Inches(12), Inches(1.2),
          ["Cùng 100 bài miniF2F, cùng budget, cùng máy với baseline no-pruning. "
           "Từ τ = 0.02 proof rate bắt đầu giảm → không dùng."], 18)

    # 3 bài toán
    s = slide(prs, "Bài toán và câu hỏi nghiên cứu", "1 · Bối cảnh")
    bullets(s, ["Prover thử nhiều tactic tại mỗi proof state; phần lớn thất bại, mỗi lần thử tốn một lời gọi Lean",
                "Dữ liệu (state, tactic, kết quả, thời gian) tăng theo định lý × nhánh → cần pipeline dữ liệu lớn",
                ("RQ1: đặc trưng proof state dự đoán tactic tiếp theo tốt đến đâu (top-k, theo domain)?", True),
                ("RQ2: pruning giảm bao nhiêu states expanded trước khi proof rate giảm?", True),
                "Primary metrics: proof success rate (chất lượng) · states expanded / bài (chi phí)"])

    # 4 kiến trúc
    s = slide(prs, "Kiến trúc 7 khối", "2 · Hệ thống")
    blocks = ["Trace\ncollector", "Event\nschema", "Parquet\nlake", "Spark\naggregation", "Features", "Predictor\n/ policy",
              "Replay\n/ eval"]
    for i, b in enumerate(blocks):
        x = Inches(0.6 + i * 1.78)
        box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.0), Inches(1.55), Inches(1.3))
        box.fill.solid()
        box.fill.fore_color.rgb = NAVY if i in (2, 3) else LIGHT
        box.line.fill.background()
        _text(s, x, Inches(2.2), Inches(1.55), Inches(1.0), b.split("\n"), 15,
              RGBColor(255, 255, 255) if i in (2, 3) else NAVY, True, PP_ALIGN.CENTER)
        if i < len(blocks) - 1:
            _text(s, x + Inches(1.5), Inches(2.4), Inches(0.35), Inches(0.5), ["→"], 20, GREY)
    bullets(s, ["Nguồn: LeanDojo Benchmark 4 (213K bước, 98K định lý) + trace tự sinh trên miniF2F qua LeanInteract",
                "Một lake duy nhất, partition source / budget_id, zstd, manifest sha256",
                "Mỗi run ghi log Phụ lục B (git commit, data version, máy, phần mềm, seed, budget) → runs.jsonl",
                "Mọi hình/bảng sinh bằng script từ results/raw; chạy hoàn toàn trong Docker"],
            y=Inches(3.9), h=Inches(3), size=18)

    # 5 schema
    s = slide(prs, "Event schema v0.1.0 và trace lake", "2 · Hệ thống")
    table(s, ["Nhóm cột", "Cột"], [
        ["Định danh", "schema_version, source, run_id, data_version, split, theorem_id, domain"],
        ["Cây tìm kiếm", "step_idx, parent_step_idx, depth, state_before, state_after"],
        ["Hành động", "tactic, tactic_rank, score, outcome, elapsed_ms"],
        ["Budget & policy", "budget_id, max_calls, timeout_s, max_depth, policy, tau"]],
        Inches(0.6), Inches(1.8), Inches(12), size=15)
    bullets(s, ["Grain = 1 lần thử tactic · khóa (run_id, theorem_id, step_idx) · parent_step_idx dựng lại cây",
                "Mỗi event mang budget của nó → không thể so sánh nhầm hai run khác budget",
                "694 MB JSON → ~18 MB Parquet; nhóm khác đọc bằng pandas/Spark + kiểm manifest"],
            y=Inches(4.3), h=Inches(2.5), size=18)

    # 6 E1
    s = slide(prs, "E1 · Spark scale: phẳng ở dữ liệu thật, tuyến tính khi nhân bản", "3 · Kết quả")
    picture(s, f"{FIG}/scale_aggregate.png", Inches(0.5), Inches(1.6), h=Inches(5.0))
    sc = scale[scale.source == "benchmark"].sort_values("events")
    real, rep = sc[sc.replication == 1], sc[sc.replication > 1]
    bullets(s, [f"Dữ liệu thật: {real.runtime_median.min():.0f}–{real.runtime_median.max():.0f} s từ "
                f"{int(real.events.min()):,} đến {int(real.events.max()):,} event → chi phí cố định lấn át",
                *[f"×{int(r.replication)}: {int(r.events) / 1e6:.1f}M event, {r.runtime_median:.0f} s, "
                  f"{r.events_per_s_median / 1e3:.0f}K event/s" for r in rep.itertuples()],
                "Median 3 lần, thanh = min–max",
                "Nhân bản chỉ để đo runtime, không dùng cho chất lượng"],
            x=Inches(7.6), y=Inches(1.8), w=Inches(5.4), size=16)

    # 7 physical plan
    s = slide(prs, "Spark: join nào gây shuffle?", "3 · Kết quả")
    bullets(s, ["Đọc lake: partition pruning (source=benchmark) + column pruning (13 cột, bỏ text state)",
                "Lọc theo tập định lý: broadcast join – bảng keep nhỏ gửi tới mọi executor, không shuffle bảng lớn",
                "transitions: self-join cha → con; phía cha nhỏ → BroadcastHashJoin, cả hai lớn → Exchange + "
                "SortMergeJoin",
                "position: Window lag theo (run_id, theorem_id) → shuffle theo khóa phân vùng",
                "Plan đầy đủ: results/raw/spark_plans/<mức>/<bảng>.txt (aggregate --explain)"], size=19)

    # hành vi tìm kiếm
    s = slide(prs, "Spark trên trace search: chi phí nằm ở đâu?", "3 · Kết quả")
    picture(s, f"{FIG}/search_behavior.png", Inches(0.4), Inches(1.6), w=Inches(8.2))
    oc = pd.read_csv(f"{RAW}/agg/search_nall/outcomes.csv").set_index("tactic_head")
    cols = [c for c in ("finished", "progress", "duplicate", "error") if c in oc]
    sh = oc[cols].div(oc[cols].sum(axis=1), axis=0)
    dup = sh[sh.duplicate > 0.5].duplicate.sort_values(ascending=False)
    br = pd.read_csv(f"{RAW}/agg/search_nall/branching.csv")
    slow = pd.read_csv(f"{RAW}/agg/search_nall/tactic_stats.csv").set_index("tactic_head").avg_elapsed_ms.sort_values()
    bullets(s, [f"Error chiếm {oc.error.sum() / oc[cols].to_numpy().sum():.0%} lời gọi",
                *[f"{t}: {v:.0%} lời gọi trả về state đã gặp (no-op)" for t, v in dup.items()],
                f"Mỗi state thử ~{br.avg_tried.median():.0f} tactic, chỉ ~{br.avg_branching.median():.1f} con mới "
                "→ cây hẹp",
                f"{slow.index[-1]} {slow.iloc[-1] / 1000:.1f} s/lời gọi, {slow.index[0]} {slow.iloc[0]:.0f} ms "
                "→ budget tính bằng giây Lean"],
            x=Inches(8.8), y=Inches(1.8), w=Inches(4.3), size=15)

    # 8 E2
    s = slide(prs, "E2 · Dự đoán tactic", "3 · Kết quả")
    picture(s, f"{FIG}/predictor_topk.png", Inches(0.5), Inches(1.6), h=Inches(5.0))
    mf, tf = pm["most_frequent"], pm["tfidf_sgd"]
    bullets(s, [f"Top-1: {mf['top1']:.3f} → {tf['top1']:.3f} (+{(tf['top1'] / mf['top1'] - 1) * 100:.0f}%)",
                f"Top-10: {mf['top10']:.3f} → {tf['top10']:.3f}",
                f"Test {tf['n_test']:,} bước, split theo định lý",
                f"Bỏ bigram: top-1 chỉ giảm {tf['top1'] - pm_uni['tfidf_sgd']['top1']:.3f}",
                f"Chi phí train: {tf['fit_s']:.0f} s vs {mf['fit_s']:.1f} s"],
            x=Inches(7.6), y=Inches(1.8), w=Inches(5.4), size=17)

    # 9 E3
    s = slide(prs, "E3 · Tăng budget không chứng minh thêm bài", "3 · Kết quả")

    def cell(r, m, f="{:.1f}"):
        lo, hi = r[f"{m}_min"], r[f"{m}_max"]
        med = f.format(r[f"{m}_median"])
        return med if hi - lo < 1e-9 else f"{med} [{f.format(lo)}–{f.format(hi)}]"
    rows = [[f"Baseline {b}", int(summ.loc[f'repl-baseline-{b}'].repeats),
             f"{summ.loc[f'repl-baseline-{b}'].proof_success_rate_median:.2f}",
             cell(summ.loc[f"repl-baseline-{b}"], "avg_states_expanded"),
             cell(summ.loc[f"repl-baseline-{b}"], "avg_wall_s")] for b in "SML"]
    table(s, ["Run (100 bài miniF2F)", "Lần", "Proof rate", "States / bài", "Wall s / bài"], rows,
          Inches(0.6), Inches(1.8), Inches(12), size=16)
    bullets(s, ["Chi phí tăng ~3× từ S lên L, proof rate đứng yên",
                "Giới hạn chất lượng nằm ở không gian hành động (14 tactic cố định, không tham số)",
                "Chi phí nằm ở các nhánh vô ích → chỗ để pruning"], y=Inches(4.2), h=Inches(2.5), size=18)

    # 10 E4 pareto
    s = slide(prs, f"E4 · Pruning: τ = {FINAL_TAU:g} cắt {save * 100:.0f}% chi phí, không mất bài", "3 · Kết quả")
    picture(s, f"{FIG}/pareto_{m_run}.png", Inches(0.5), Inches(1.6), h=Inches(5.0))
    t0005 = summ.loc["repl-prune-M-tau0005"]
    swt = sw.set_index("tau")
    bullets(s, [f"Baseline M: {bM.proof_success_rate_median:.2f} · {bM.avg_states_expanded_median:.1f} states",
                f"τ=0.005 live: {t0005.proof_success_rate_median:.2f} · {t0005.avg_states_expanded_median:.1f}",
                f"τ=0.01 live: {fin.proof_success_rate_median:.2f} · {fin.avg_states_expanded_median:.1f}",
                f"Replay τ=0.01 dự đoán {swt.loc[0.01].proof_success_rate:.2f} (unknown = thất bại → bảo thủ)",
                f"Điểm gối: τ=0.02 → {swt.loc[0.02].proof_success_rate:.2f}, τ=0.05 → "
                f"{swt.loc[0.05].proof_success_rate:.2f}"],
            x=Inches(7.6), y=Inches(1.8), w=Inches(5.4), size=16)

    # 11 losses
    s = slide(prs, "Vì sao τ cao làm mất bài?", "3 · Kết quả")
    table(s, ["τ"] + list(cause.columns), [[f"{t:g}"] + [int(v) for v in r] for t, r in cause.iterrows()],
          Inches(0.6), Inches(1.8), Inches(12), size=15)
    losses = pd.read_csv(f"{FIG}/pruning_losses_M.csv")
    cut = losses[losses.tau == 0.05].pruned_proof_tactics.str.split(" ; ").explode().value_counts()
    l01 = losses[losses.tau == 0.01]
    rescued = sorted(set(live_rescue(l01.theorem_id.tolist()).values()))
    bullets(s, ["Mọi bài mất đều do policy cắt đúng tactic trên đường chứng minh gốc (τ=0.05: "
                + ", ".join(f"{k} ×{v}" for k, v in cut.items()) + ")",
                f"Mathlib: rw/simp/exact = {share(['rw', 'simp', 'exact']):.0%} số bước; rfl {share(['rfl']):.1%}, "
                f"linarith {share(['linarith']):.1%} → prior thấp, dưới ngưỡng",
                f"τ=0.01: replay mất {len(l01)} bài, live vẫn chứng minh bằng {', '.join(rescued)} (ngoài trace)",
                "Hướng cải tiến: fine-tune predictor trên trace search của chính prover"],
            y=Inches(4.0), h=Inches(3), size=17)

    # 12 giới hạn
    s = slide(prs, "Giới hạn và quyết định thiết kế", "4 · Thảo luận")
    bullets(s, ["Search live 100 bài (Lean ~20 s/bài); quy mô lớn ở analytics + lake nhân bản",
                "Replay là ước lượng bảo thủ → τ cuối cùng kiểm chứng bằng chạy live, cùng budget",
                "Deviation: LeanInteract thay LeanDojo cho search (không phải trace mathlib) – cùng schema, ghi rõ",
                "Spark local[8] một node; kết quả scale = song song trong một máy",
                "Lỗi REPL timeout được phát hiện, sửa có test; 5 run cách ly kèm lý do, chạy lại toàn bộ"], size=19)

    # 13 demo
    s = slide(prs, "Demo", "5 · Demo")
    bullets(s, ["make test → 27 test xanh",
                "Mở lake: pd.read_parquet(..., filters=[('source','=','search')]) → xem 1 cây tìm kiếm",
                "make aggregate-search → bảng tactic_stats, branching; mở Spark UI :4040",
                "make dashboard → http://localhost:8501 (5 tab)",
                "Pareto + bảng bài bị mất trong results/figures/"], size=20)
    os.makedirs(os.path.dirname(out), exist_ok=True)
    prs.save(out)
    print("->", out, len(prs.slides), "slides")


if __name__ == "__main__":
    build()
